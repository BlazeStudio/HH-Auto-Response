"""Excel-отчёт по откликам (openpyxl): сводка + листы по результату, со ссылками на вакансии."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .results import APPLIED, FAILED, HH_DAILY_LIMIT, QUESTIONS, SKIPPED, Summary, response_form_url

HEADER_FILL = PatternFill("solid", fgColor="1F4E9E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def build_workbook(records: list[dict], summary: Summary) -> Workbook:
    applied_ids = {r["id"] for r in records if r["category"] == APPLIED}
    wb = Workbook()
    _summary_sheet(wb.active, summary)

    applied = [r for r in records if r["category"] == APPLIED]
    _table(wb.create_sheet("Откликнулся"), applied, [
        ("Дата", 19, lambda r: r["at"]),
        ("Вакансия", 50, lambda r: (r["title"], r["url"])),
        ("Компания", 34, lambda r: r["company"]),
        ("Письмо", 9, lambda r: "да" if r.get("letter") else "нет"),
        ("Комментарий", 50, lambda r: r.get("reason", "")),
    ])
    # По пропущенным и неудачным — последняя попытка по каждой вакансии, если отклик так и не ушёл
    pending = _latest_per_vacancy([r for r in records if r["id"] not in applied_ids])
    _table(wb.create_sheet("Вопросы"), [r for r in pending if r["category"] == QUESTIONS], [
        ("Дата", 19, lambda r: r["at"]),
        ("Вакансия", 42, lambda r: (r["title"], r["url"])),
        ("Компания", 30, lambda r: r["company"]),
        ("Вопросы работодателя", 70, lambda r: "\n".join(f"{i}. {q}" for i, q in enumerate(r.get("questions"), 1))
         or r.get("reason", "")),
        ("Ответить и откликнуться", 26, lambda r: ("Открыть форму отклика", response_form_url(r["id"]))),
    ], note="Вакансии, где работодатель задал вопросы. Откройте форму отклика по ссылке и ответьте сами.")
    _table(wb.create_sheet("Пропущенные"), [r for r in pending if r["category"] == SKIPPED], [
        ("Дата", 19, lambda r: r["at"]),
        ("Вакансия", 50, lambda r: (r["title"], r["url"])),
        ("Компания", 34, lambda r: r["company"]),
        ("Причина", 70, lambda r: r.get("reason", "")),
    ])
    _table(wb.create_sheet("Не откликнулся"), [r for r in pending if r["category"] == FAILED], [
        ("Дата", 19, lambda r: r["at"]),
        ("Вакансия", 50, lambda r: (r["title"], r["url"])),
        ("Компания", 34, lambda r: r["company"]),
        ("Ошибка", 70, lambda r: r.get("reason", "")),
    ], note="Ошибки и лимит hh. Такие вакансии программа попробует снова при следующем запуске.")
    return wb


def _latest_per_vacancy(records: list[dict]) -> list[dict]:
    latest: dict[str, dict] = {}
    for record in records:
        latest[record["id"]] = record  # записи идут по времени — остаётся последняя
    return sorted(latest.values(), key=lambda r: r["at"], reverse=True)


def _table(ws: Worksheet, rows: list[dict], columns: list, note: str = "") -> None:
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="555555")
        start = 3
    for col, (title, width, _) in enumerate(columns, 1):
        cell = ws.cell(row=start, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(col)].width = width
    for row_index, record in enumerate(sorted(rows, key=lambda r: r["at"], reverse=True), start + 1):
        for col, (_, _, getter) in enumerate(columns, 1):
            value = getter(record)
            cell = ws.cell(row=row_index, column=col)
            if isinstance(value, tuple):  # (текст, ссылка)
                cell.value, cell.hyperlink, cell.font = value[0], value[1], LINK_FONT
            else:
                cell.value = value
            cell.alignment = WRAP if isinstance(value, str) and "\n" in value else TOP
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(columns))}{start + len(rows)}"


def _summary_sheet(ws: Worksheet, summary: Summary) -> None:
    ws.title = "Сводка"
    ws["A1"] = "Статистика откликов"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Лимит hh — около {HH_DAILY_LIMIT} откликов в сутки. Сегодня отправлено: {summary.today.applied}"

    headers = ["", "Сегодня", f"Последний запуск ({summary.last_run_at or '—'})", "Всего"]
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    rows = [
        ("Откликнулся", "applied"), ("  из них с письмом", "letters"), ("Вопросы работодателя", "questions"),
        ("Пропущено", "skipped"), ("Не откликнулся (ошибки, лимит)", "failed"),
    ]
    for i, (label, attr) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=label).font = BOLD if attr == "applied" else Font()
        for col, totals in enumerate((summary.today, summary.last_run, summary.total), 2):
            ws.cell(row=i, column=col, value=getattr(totals, attr))

    start = 5 + len(rows) + 2
    ws.cell(row=start - 1, column=1, value="По дням").font = BOLD
    day_headers = ["Дата", "Откликнулся", "С письмом", "Вопросы", "Пропущено", "Не откликнулся"]
    for col, title in enumerate(day_headers, 1):
        cell = ws.cell(row=start, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, (day, t) in enumerate(summary.by_day.items(), start + 1):
        for col, value in enumerate((day, t.applied, t.letters, t.questions, t.skipped, t.failed), 1):
            ws.cell(row=i, column=col, value=value)
    for col, width in zip("ABCDEF", (32, 14, 34, 12, 12, 16)):
        ws.column_dimensions[col].width = width

